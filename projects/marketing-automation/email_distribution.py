"""Sanitized portfolio example of an automated reporting workflow.

This version uses fictional file names, generic field names, and no real recipients.
"""

from __future__ import annotations

import html
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

try:
    import win32com.client as win32
except ImportError:
    win32 = None

BASE_DIR = Path(__file__).resolve().parent
REPORT_BOOK = BASE_DIR / "demo_location_reports.xlsx"
RECIPIENT_MAP = BASE_DIR / "demo_recipients.csv"


def load_recipients() -> dict[str, str]:
    recipients = pd.read_csv(RECIPIENT_MAP, dtype=str).fillna("")
    recipients.columns = recipients.columns.str.strip()

    required = {"Location", "Email"}
    missing = required.difference(recipients.columns)
    if missing:
        raise ValueError(f"Missing columns: {sorted(missing)}")

    recipients["Location"] = recipients["Location"].str.strip()
    recipients["Email"] = recipients["Email"].str.strip()
    recipients = recipients[
        (recipients["Location"] != "") & (recipients["Email"] != "")
    ]
    return dict(zip(recipients["Location"], recipients["Email"]))


def dataframe_to_html(df: pd.DataFrame) -> str:
    safe = df.fillna("").copy()
    headers = "".join(f"<th>{html.escape(str(c))}</th>" for c in safe.columns)
    rows = []
    for _, row in safe.iterrows():
        cells = "".join(f"<td>{html.escape(str(v))}</td>" for v in row)
        rows.append(f"<tr>{cells}</tr>")
    return f"<table><thead><tr>{headers}</tr></thead><tbody>{''.join(rows)}</tbody></table>"


def send_reports() -> None:
    if win32 is None:
        raise RuntimeError("Outlook COM is only available on Windows with pywin32.")

    recipients = load_recipients()
    workbook = load_workbook(REPORT_BOOK, read_only=True)
    sheet_names = workbook.sheetnames
    workbook.close()

    outlook = win32.Dispatch("Outlook.Application")

    with tempfile.TemporaryDirectory(prefix="demo_reports_") as tmp_dir:
        tmp = Path(tmp_dir)

        for location in sheet_names:
            recipient = recipients.get(location)
            if not recipient:
                continue

            df = pd.read_excel(REPORT_BOOK, sheet_name=location)
            attachment = tmp / f"{location}_report.xlsx"
            df.to_excel(attachment, index=False)

            mail = outlook.CreateItem(0)
            mail.To = recipient
            mail.Subject = f"Weekly report - {location}"
            mail.HTMLBody = (
                "<p>Hello,</p>"
                "<p>Attached is your weekly report.</p>"
                + dataframe_to_html(df)
            )
            mail.Attachments.Add(str(attachment.resolve()))
            mail.Send()


if __name__ == "__main__":
    send_reports()
